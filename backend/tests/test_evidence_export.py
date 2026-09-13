import csv
import io

from docx import Document

from app.services.evidence_export import build_csv, build_docx

ROWS = [
    {
        "question_text": "Do you log all AI agent actions?",
        "final_answer": "Yes, every action is logged automatically.",
        "draft_answer": "Draft: yes.",
        "status": "approved",
        "evidence_event_ids": ["evt_1", "evt_2"],
    },
    {
        "question_text": "Is PII redacted before storage?",
        "final_answer": None,
        "draft_answer": "Draft: PII is redacted via Presidio.",
        "status": "draft",
        "evidence_event_ids": [],
    },
]


def test_build_csv_prefers_final_answer_over_draft():
    content = build_csv(ROWS)
    rows = list(csv.reader(io.StringIO(content.decode("utf-8"))))
    assert rows[0] == ["Question", "Answer", "Status", "Evidence Event IDs"]
    assert rows[1] == [
        "Do you log all AI agent actions?",
        "Yes, every action is logged automatically.",
        "approved",
        "evt_1, evt_2",
    ]
    # falls back to the draft when there's no reviewed final answer yet
    assert rows[2] == [
        "Is PII redacted before storage?",
        "Draft: PII is redacted via Presidio.",
        "draft",
        "",
    ]


def test_build_csv_handles_no_answer_at_all():
    rows = [{"question_text": "Untouched question?", "final_answer": None, "draft_answer": None, "status": "draft", "evidence_event_ids": []}]
    content = build_csv(rows)
    parsed = list(csv.reader(io.StringIO(content.decode("utf-8"))))
    assert parsed[1][1] == ""


def test_build_docx_contains_questions_and_answers():
    content = build_docx("Acme Inc", "questionnaire.csv", ROWS, watermark=False)
    doc = Document(io.BytesIO(content))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(h.text for h in doc.paragraphs)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Acme Inc" in full_text
    assert "questionnaire.csv" in full_text
    assert any("Do you log all AI agent actions?" in p.text for p in doc.paragraphs)
    assert any("Yes, every action is logged automatically." in p.text for p in doc.paragraphs)
    assert any("evt_1, evt_2" in p.text for p in doc.paragraphs)
    assert not any("FREE PLAN" in p.text for p in doc.paragraphs)


def test_build_docx_adds_watermark_on_free_plan():
    content = build_docx("Acme Inc", "questionnaire.csv", ROWS, watermark=True)
    doc = Document(io.BytesIO(content))
    assert any("FREE PLAN" in p.text for p in doc.paragraphs)


def test_build_docx_handles_empty_rows():
    content = build_docx("Acme Inc", "questionnaire.csv", [], watermark=False)
    doc = Document(io.BytesIO(content))
    assert any("Acme Inc" in p.text for p in doc.paragraphs)
